import openpyxl

from ffs.importers.thickness_workbook import import_workbook


def _mini(path):
    wb = openpyxl.Workbook()
    idx = wb.active
    idx.title = "indexes"
    idx.append([None, "Floors", None, None, None, None, "Roofs", None, None])
    idx.append([None, "UL Designs", "N661", "N661", "N635", "N635", None, "UL Design", "Y669"])
    idx.append(
        [
            None,
            None,
            "FIRESOLVE® SBLW",
            "FIRESOLVE® SBNW",
            "SprayFilm® WB 3LW",
            "SprayFilm® WB 3NW",
            None,
            None,
            "FIRESOLVE® SB",
        ]
    )
    wide = wb.create_sheet("Unrestrain Floor Beam Thickness")
    wide.append([])
    wide.append(
        [
            None,
            None,
            None,
            None,
            None,
            "FIRESOLVE® SBLW",
            "FIRESOLVE® SBLW",
            "FIRESOLVE® SBNW",
            "FIRESOLVE® SBNW",
            "SprayFilm® WB 3LW",
            "SprayFilm® WB 3LW",
        ]
    )
    wide.append([None, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10])
    wide.append(
        [
            None,
            None,
            None,
            None,
            None,
            "FIRESOLVE® SB",
            "FIRESOLVE® SB",
            "FIRESOLVE® SB",
            "FIRESOLVE® SB",
            "SprayFilm® WB 3",
            "SprayFilm® WB 3",
        ]
    )
    wide.append([None, None, None, None, None, "LW", "LW", "NW", "NW", "LW", "LW"])
    wide.append(
        [
            None,
            "Designation",
            "Unrestrained",
            "W/D",
            "SF/LF",
            "1 Hr",
            "2 Hr",
            "1 Hr",
            "2 Hr",
            "1 Hr",
            "2 Hr",
        ]
    )
    wide.append([None, "W18 x 35", "Unrestrained", 0.83, 4.29, 41, 87, 36, 69, "NR", 102])
    wide.append([None, "W12 x 26", "Unrestrained", 0.66, 3.33, 47, 99, 40, 81, 60, "NR"])
    roof = wb.create_sheet("Roof beam Thicknesses")
    roof.append([])
    roof.append([None, None, None, None, None, "FIRESOLVE® SB", "FIRESOLVE® SB"])
    roof.append([None, "Designation", None, "W/D", "SF/LF", "1 Hr", "1-1/2 Hr"])
    roof.append([None, "W8 x 10", None, 0.37, 2.46, 55, 110])
    long = wb.create_sheet("N614 - NW")
    long.append(
        ["Designation", "Unrestrained", "Rating", "W/D", "Mils", "sq ft factor", "sq ft factor"]
    )
    long.append(["W18 x 35", "Unrestrained", 1, 0.83, 39, 4.29, 4.29])
    long.append(["W18 x 35", "Unrestrained", 1.5, 0.83, 66, 4.29, 4.29])
    long.append(["Square Tube - 6 x 6 x 1/4", "Unrestrained", 2, 0.61, "NR", 2.0, 2.0])
    wb.create_sheet("TOTAL").append(["MEMBER DESIGNATION", "TOTAL SQ. FT."])
    wb.save(path)


def test_wide_and_long(tmp_path):
    p = tmp_path / "mini.xlsx"
    _mini(p)
    imp = import_workbook(p, product="SprayFilm WB 4")
    assert imp.sheets_parsed == {
        "Unrestrain Floor Beam Thickness": "wide",
        "Roof beam Thicknesses": "wide",
        "N614 - NW": "long",
    }
    assert "TOTAL" not in imp.sheets_parsed  # summary sheets are ignored by name
    assert imp.design_index["Floors"]["firesolvesblw"] == "N661"
    wide = [r for r in imp.rows if r.sheet == "Unrestrain Floor Beam Thickness"]
    assert len(wide) == 12
    r = wide[0]
    assert (
        r.member_label,
        r.canonical,
        r.product,
        r.variant,
        r.design,
        r.rating_hr,
        r.thickness_mils,
    ) == ("W18 x 35", "W18X35", "FIRESOLVE® SB", "LW", "N661", 1.0, 41.0)
    assert (
        r.section == "Floors"
        and r.section_factor == 0.83
        and r.section_factor_kind == "W/D"
        and r.sf_per_lf == 4.29
    )
    assert r.restraint == "unrestrained" and r.row == 7 and r.col == 6
    nw = [x for x in wide if x.variant == "NW"]
    assert nw[0].thickness_mils == 36.0 and nw[0].design == "N661"
    wb3 = [x for x in wide if x.product == "SprayFilm® WB 3"]
    assert wb3[0].not_rated and wb3[0].thickness_mils is None and wb3[0].design == "N635"
    roof = [r for r in imp.rows if r.sheet == "Roof beam Thicknesses"]
    assert [x.rating_hr for x in roof] == [1.0, 1.5] and roof[0].design == "Y669"
    long = [r for r in imp.rows if r.sheet == "N614 - NW"]
    assert (
        len(long) == 3
        and long[0].design == "N614"
        and long[0].variant == "NW"
        and long[0].product == "SprayFilm WB 4"
    )
    assert long[1].rating_hr == 1.5 and long[1].thickness_mils == 66.0
    assert (
        long[2].not_rated and long[2].canonical is None
    )  # "Square Tube - 6 x 6 x 1/4" is not an AISC label
    assert long[2].section_factor == 0.61
