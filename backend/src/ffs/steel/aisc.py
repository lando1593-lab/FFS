"""Authoritative shape validation against the AISC Shapes Database.

AISC publishes the Shapes Database (v16.0 at time of writing) as an Excel workbook on aisc.org.
It is not redistributed in this repository. Download it locally and point ``FFS_AISC_SHAPES`` (or
``--aisc``) at the ``.xlsx``/``.csv``. When absent, designations are syntactically parsed but
``validated`` is False and identification confidence is capped.

Columns used (v16 naming): ``AISC_Manual_Label`` (e.g. ``W18X35``), ``W`` (plf), ``d``, ``bf``,
``tw``, ``tf``, ``A``. Section-factor math (Milestone 3) will use the geometric columns.
"""

from __future__ import annotations

import csv
import os
from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class SteelShape:
    label: str  # canonical AISC label
    shape_type: str  # W, HSS, L, ...
    weight_plf: float | None
    depth_in: float | None
    bf_in: float | None
    tw_in: float | None
    tf_in: float | None
    area_in2: float | None
    source: str  # file name + version


class ShapeDatabase:
    def __init__(self, shapes: dict[str, SteelShape], source: str):
        self._shapes = shapes
        self.source = source

    @classmethod
    def load(cls, path: str | Path | None = None) -> ShapeDatabase | None:
        raw = str(path) if path else os.environ.get("FFS_AISC_SHAPES", "")
        if not raw:
            return None
        p = Path(raw)
        if not p.is_file():
            return None
        rows: list[dict[str, str]]
        if p.suffix.lower() == ".csv":
            with p.open(newline="", encoding="utf-8-sig") as f:
                rows = list(csv.DictReader(f))
        else:
            import openpyxl

            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            it = ws.iter_rows(values_only=True)
            header = [str(h).strip() if h is not None else "" for h in next(it)]
            rows = [
                dict(zip(header, [("" if v is None else str(v)) for v in r], strict=False))
                for r in it
            ]
        shapes: dict[str, SteelShape] = {}

        def num(v: str) -> float | None:
            v = (v or "").strip()
            if v in ("", "–", "-", "—"):
                return None
            try:
                return float(v)
            except ValueError:
                return None

        for r in rows:
            label = (r.get("AISC_Manual_Label") or r.get("Label") or "").strip().upper()
            if not label:
                continue
            shapes[label] = SteelShape(
                label=label,
                shape_type=(r.get("Type") or label.rstrip("0123456789.X/-"))[:10],
                weight_plf=num(r.get("W", "")),
                depth_in=num(r.get("d", "")),
                bf_in=num(r.get("bf", "")),
                tw_in=num(r.get("tw", "")),
                tf_in=num(r.get("tf", "")),
                area_in2=num(r.get("A", "")),
                source=p.name,
            )
        return cls(shapes, p.name)

    def get(self, canonical: str) -> SteelShape | None:
        return self._shapes.get(canonical.upper())

    def __contains__(self, canonical: str) -> bool:
        return canonical.upper() in self._shapes

    def __len__(self) -> int:
        return len(self._shapes)
