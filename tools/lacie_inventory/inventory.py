#!/usr/bin/env python3
"""LaCie reference-library inventory — READ ONLY.

Walks a library root, opens every file read-only, and writes a catalog OUTSIDE the library:
  <out>/catalog.jsonl    one JSON object per file, appended as the walk proceeds (survives Ctrl-C)
  <out>/catalog.sqlite   full catalog (one row per file + classification + extracted hints)
  <out>/catalog.csv      flat summary for Excel (UTF-8 with BOM)
  <out>/summary.md       counts only (document types, manufacturers, decades, scanned PDFs,
                         error counts) — safe to share
  <out>/summary_paths.md top folders, error paths, unclassified paths — contains file/folder
                         names, keep local
  <out>/run.log          what happened (start/stop, interrupted?, totals)

Safety:
  * Never writes inside the library root, and refuses an --out on the same volume as the root.
  * Opens files read-only by path; never renames, moves, or deletes. Reading can still update
    "last accessed" timestamps on some filesystems; contents and names are never touched.
  * Unreadable directories, bad timestamps, corrupt/encrypted files are recorded as error rows;
    nothing stops the walk. Ctrl-C writes a partial catalog marked PARTIAL.
  * Symlinks and Windows junctions are never followed.
  * Stores at most a short text snippet per file (default 600 chars) — no full-text copies.
  * Photos, video and audio are never opened or hashed: by default they are only counted per
    top-level folder (--media count); --media list records their names and sizes instead.
    --exclude "Pattern" skips whole folders (e.g. personal folders) by name.

Classification is heuristic keyword scoring with a confidence in [0, 1]. It is a *map* of the
library for humans to review; it is never an authority on current code or product compliance
(ADR-0004).

Usage:
  python inventory.py --find                       # list candidate drives / mount points
  python inventory.py /Volumes/LaCie --out ../data/lacie_catalog
  python inventory.py D:\\ --out ..\\data\\lacie_catalog
  python inventory.py <root> --max-files 500 --dry-run
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sqlite3
import stat
import sys
import time
from dataclasses import asdict, dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path

# ---------------------------------------------------------------------------------------------
# Domain vocab (keywords only; nothing here is a thickness/design value)
# ---------------------------------------------------------------------------------------------
MANUFACTURERS = {
    "isolatek": ["isolatek", "cafco", "blaze-shield", "blaze shield", "blazeshield", "fendolite", "sprayfilm"],
    "gcp/grace": ["monokote", "mk-6", "mk6", "mk 6", "z-146", "z146", "z-156", "z156", "z-106", "z106",
                  "w.r. grace", "w. r. grace", "wr grace", "grace construction", "gcp applied", "gcp "],
    "carboline": ["carboline", "pyrocrete", "pyrolite", "thermo-sorb", "thermosorb", "thermo-lag", "thermolag",
                  "pyroclad", "southwest type", "southwest 5gp", "southwest 7gp"],
    "sherwin-williams": ["firetex", "sherwin"],
    "ppg": ["pitt-char", "pittchar", "steelguard", "ppg"],
    "akzonobel/international": ["interchar", "chartek", "akzo", "international paint", "international protective"],
    "hilti": ["hilti", "cfp-sp"],
    "nullifire": ["nullifire"],
    "promat": ["promat"],
    "hempel": ["hempel"],
    "jotun": ["jotun", "steelmaster"],
}
DOC_TYPES = {
    "ul_design": ["bxuv", "design no", "design number", "fire resistance rating", "restrained assembly rating",
                  "unrestrained beam rating", "unrestrained assembly rating", "ul 263", "ul263", "astm e119", "e-119",
                  "product iq", "classified by underwriters"],
    "listing_report": ["evaluation report", "esr-", "icc-es", "icc es", "intertek", "warnock hersey", "certificate of",
                       "certification", "listing report", "design listing", "fm approvals", "test report"],
    "shape_table": ["w/d", "weight per foot", "heated perimeter", "section factor", "nominal depth", "flange width",
                    "web thickness", "flange thickness", "wt/ft", "lbs/ft", "perimeter"],
    "thickness_chart": ["thickness chart", "thickness schedule", "thickness table", "required thickness",
                        "minimum thickness", "thickness required", "hourly rating", "hour rating"],
    "spray_chart": ["spray chart", "color up", "color-up", "colorup", "spray schedule", "spray report",
                    "fireproofing schedule", "spray items"],
    "shop_drawing": ["shop drawing", "fireproofing plan", "fp-1", "fp-2", "fp1.", "fp2.", "key plan", "fireproofing legend"],
    "submittal": ["submittal", "transmittal", "approved as noted", "revise and resubmit", "no exceptions taken",
                  "reviewed - no exception", "furnish as corrected", "make corrections noted"],
    "product_data": ["product data sheet", "technical data sheet", "data sheet", "physical properties", "packaging",
                     "shelf life", "coverage rate", "typical properties", "product description", "basic uses",
                     "limitations", "storage"],
    "sds": ["safety data sheet", "material safety data", "msds", "hazard identification", "ghs", "first aid"],
    "primer_topcoat_list": ["primer list", "approved primers", "acceptable primers", "primer guide", "topcoat list",
                            "acceptable topcoat", "approved topcoat", "topcoat guide", "primer compatibility",
                            "compatible primers", "primers for"],
    "technical_bulletin": ["technical bulletin", "tech bulletin", "tech release", "technical release", "technical note",
                           "technical letter", "bulletin"],
    "specification": ["section 07 81", "07 81 00", "07 81 16", "07 81 23", "078100", "078116", "078123", "07810",
                      "07250", "sprayed fire-resistive", "spray-applied fire-resistive", "part 1 - general",
                      "part 2 - products", "part 3 - execution", "guide specification"],
    "estimate": ["takeoff", "take-off", "bid recap", "recap", "board feet", "bd ft", "bd. ft.", "bags required",
                 "unit price", "quantity survey", "bid form", "estimating program", "estimating guide",
                 "estimating worksheet", "budget worksheet"],
    "proposal_bid_letter": ["proposal", "bid letter", "we propose", "scope of work", "exclusions", "inclusions",
                            "bid qualification", "qualifications and exclusions", "lump sum", "price includes"],
    "labor_pricing": ["labor matrix", "labor rate", "rate sheet", "per hour", "hourly rate", "wage", "burden",
                      "production rate", "crew size", "man hours", "manhours", "sq ft per day", "ocip", "ccip",
                      "insurance calculation"],
    "procedure": ["guidelines and procedures", "estimating guidelines", "step by step", "start to finish",
                  "checklist", "cheat sheet", "quick reference", "how to", "standard operating"],
    "application_guide": ["application guide", "application instructions", "application manual",
                          "installation manual", "applicator", "mixing instructions", "pump", "nozzle",
                          "wet mix", "dry mix", "surface preparation"],
    "yield_chart": ["yield", "coverage chart", "bd ft/bag", "board feet per bag", "sq ft per gallon",
                    "theoretical coverage"],
    "field_drawing": ["field drawing", "field set", "no spray", "do not spray", "field use", "field copy"],
    "inspection": ["inspection report", "special inspection", "astm e605", "e605", "density test results",
                   "bond test results", "thickness readings", "gauge readings", "inspector"],
    "structural_drawing": ["framing plan", "beam schedule", "column schedule", "structural notes", "foundation plan",
                           "typical details"],
    "architectural_drawing": ["floor plan", "reflected ceiling", "wall section", "finish schedule", "code summary",
                              "code analysis", "life safety plan", "building section"],
    "code_reference": ["international building code", "ibc 20", "ibc section", "section 704", "section 703",
                       "table 601", "significant changes", "code change", "nfca"],
    "sustainability_compliance": ["leed", "voc content", "hpd", "chpd", "epd", "environmental product declaration",
                                  "cdph", "low emitting", "sustainability"],
    "warranty": ["warranty", "warrants", "warranted"],
    "contact_list": ["contact list", "phone", "email", "rep ", "sales representative", "territory manager"],
    "correspondence": ["dear ", "regards,", "sincerely", "rfi", "request for information", "change order", "memo",
                       "letter"],
}
# Strong filename cues (regex → (doc_type, bonus)). Filenames are the most reliable signal in a
# human-organised library; a hit here outweighs a few body keywords.
FILENAME_RULES = [
    (r"\b(?:c-|i-)?tds\b|\bpds\b|data[ _-]?sheet|gcpat_|monokote_|_us_\d{4}", "product_data", 12),
    (r"\bsds\b|\bmsds\b", "sds", 10),
    (r"primer|topcoat|top[ _-]coat", "primer_topcoat_list", 7),
    (r"tech(?:nical)?[ _-]?(?:release|bulletin|note)", "technical_bulletin", 8),
    (r"estimating[ _-]?(?:guide|program|worksheet)|takeoff|take-off|budget[ _-]worksheet", "estimate", 7),
    (r"labor[ _-]?matrix|rate[ _-]?sheet|ocip|ccip", "labor_pricing", 8),
    (r"bid[ _-]?letter|proposal|qualification", "proposal_bid_letter", 8),
    (r"guidelines|procedures|cheat[ _-]?sheet|quick[ _-]?reference|start[ _-]to[ _-]finish|estimating[ _-]a[ _-]job",
     "procedure", 7),
    (r"warranty", "warranty", 10),
    (r"contact[ _-]?list", "contact_list", 10),
    (r"evaluation|\besr\b|icc|certifica", "listing_report", 7),
    (r"^(?:bxuv[._-]?)?[dnpsxy]-?\d{3,4}[a-z]?(?:\s*\(dup\))?\.pdf$", "ul_design", 12),
    (r"wide[ _-]?flange|angles?|channels?|tube[ _-]?steel|\bhss\b|pipes?|round[ _-]?bars?|wt[ _-]?columns?|beams?\.pdf|columns?\.pdf",
     "shape_table", 8),
    (r"spray[ _-]?(?:chart|report)|shop[ _-]?print|color[ _-]?up", "spray_chart", 8),
    (r"shop[ _-]?drawing", "shop_drawing", 8),
    (r"application[ _-]?(?:guide|manual)|installation[ _-]?manual", "application_guide", 8),
    (r"leed|voc|chpd|hpd", "sustainability_compliance", 7),
    (r"sig(?:nificant)?[ _-]?changes|ibc|nfca|\b70[34]\.", "code_reference", 8),
    (r"awarded[ _-]jobs|task[ _-]tracker|company[ _-]info", "company_admin", 10),
]
FILENAME_RULES = [(re.compile(rx, re.IGNORECASE), dt, bonus) for rx, dt, bonus in FILENAME_RULES]

RATING_RE = re.compile(
    r"(?<![\w/])(?P<r>\d(?:\s*-\s*\d/\d|\s+\d/\d)?|\d/\d|\d\.5)\s*-?\s*(?:hr|hrs|hour|hours)\b", re.IGNORECASE
)
UL_STRICT_RE = re.compile(
    r"(?:BXUV\s*\.?\s*|(?:UL\s+)?design\s*(?:no\.?|number|#)?\s*:?\s*)(?P<d>[A-Z]{1,3}-?\d{3,4}[A-Z]?)\b",
    re.IGNORECASE,
)
UL_LOOSE_RE = re.compile(r"(?<![A-Za-z0-9/.-])(?P<d>[DNPSXY]-?\d{3,4})(?![A-Za-z0-9.-])")
INTERTEK_RE = re.compile(r"\b(?:CC|FP|IF|WH)\s*/\s*(?:IF|FP|CC)?\s*\d{2,4}\s*-\s*\d{1,3}\b|\bWHI[- ]?\d{3,5}\b", re.IGNORECASE)
YEAR_RE = re.compile(r"(?<![\d.-])(19[5-9]\d|20[0-4]\d)(?![\d.-])(?!\s*(?:psi|psf|pcf|°|deg|lb|lbs|kg|sf|sq|bd|ft|mils?|gal|cfm|rpm))", re.IGNORECASE)
DATE_RE = re.compile(
    r"\b(\d{1,2}[/-]\d{1,2}[/-](?:19|20)?\d{2})\b|\b((?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+\d{1,2},?\s+(?:19|20)\d{2}|(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+(?:19|20)\d{2})\b",
    re.IGNORECASE,
)
MEMBER_RE = re.compile(
    r"(?<![A-Z0-9])(W\d{1,2}X\d{1,3}|HP\d{1,2}X\d{1,3}|HSS\d+(?:\.\d+)?X\d+(?:\.\d+)?(?:X\d+/\d+)?|WT\d{1,2}X\d{1,3}|"
    r"C\d{1,2}X\d{1,2}(?:\.\d)?|MC\d{1,2}X\d{1,3}|L\d(?:-\d/\d)?X\d(?:-\d/\d)?X\d/\d{1,2}|PIPE\s*\d+|"
    r"\d{2}(?:K|LH|DLH)\d{1,2})(?![A-Z0-9])",
    re.IGNORECASE,
)
MEMBER_WORDS_RE = re.compile(r"\b(joist|girder|column|beam|deck|truss|brace|lintel|purlin)s?\b", re.IGNORECASE)
FP_TYPE = {
    "sfrm": ["sfrm", "spray-applied", "spray applied", "sprayed fire", "cementitious", "fibrous", "monokote", "cafco",
             "blaze-shield", "pyrocrete", "pyrolite", "z-146", "mk-6", "wet mix", "dry mix", "bags"],
    "intumescent": ["intumescent", "thin film", "thin-film", "dft", "mils", "firetex", "interchar", "thermo-sorb",
                    "sprayfilm", "pyroclad", "chartek", "steelguard", "pitt-char", "epoxy"],
    "board/other": ["board fireproofing", "fire board", "mineral wool", "gypsum wrap", "membrane protection",
                    "fire-rated board", "ceramic blanket"],
}
PROJECT_HINT_RE = re.compile(
    r"\b(?:project|job)\s*(?:name|no\.?|number|#|title)\s*[:\-#]?\s*(?P<p>[A-Za-z0-9][^\n]{3,60})|"
    r"\b(?:project|job)\s*:\s*(?P<q>[A-Za-z0-9][^\n]{3,60})",
    re.IGNORECASE,
)

TEXT_EXT = {".txt", ".csv", ".md", ".rtf"}
PDF_EXT = {".pdf"}
XLS_EXT = {".xlsx", ".xlsm", ".xls"}
DOC_EXT = {".docx", ".doc"}
IMG_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".gif", ".heic", ".heif", ".raw", ".cr2", ".nef",
           ".dng", ".webp", ".psd"}
VIDEO_EXT = {".mp4", ".mov", ".avi", ".m4v", ".mkv", ".wmv", ".mts", ".m2ts", ".3gp", ".mpg", ".mpeg"}
AUDIO_EXT = {".mp3", ".m4a", ".wav", ".aac", ".flac", ".wma"}
ARCHIVE_EXT = {".zip", ".7z", ".rar", ".tar", ".gz", ".dmg", ".iso"}
MEDIA_CLASSES = {"image", "video", "audio"}
CAD_EXT = {".dwg", ".dxf", ".dwf", ".rvt", ".ifc", ".skp"}
SKIP_NAMES = {".DS_Store", "Thumbs.db", "desktop.ini", ".Spotlight-V100", ".Trashes", ".fseventsd",
              "$RECYCLE.BIN", "System Volume Information", ".TemporaryItems", ".DocumentRevisions-V100"}

WIN_REPARSE = 0x400  # FILE_ATTRIBUTE_REPARSE_POINT


@dataclass
class Entry:
    path: str
    rel_path: str
    folder: str
    filename: str
    ext: str
    size_bytes: int
    mtime_iso: str
    sha256: str | None
    file_class: str  # pdf / excel / word / image / cad / text / other / directory
    pages: int | None = None
    has_text_layer: bool | None = None
    likely_scanned: bool | None = None
    title_meta: str | None = None
    author_meta: str | None = None
    created_meta: str | None = None
    snippet: str = ""
    doc_type: str = "unclassified"
    doc_type_confidence: float = 0.0
    doc_type_scores: dict[str, int] = field(default_factory=dict)
    manufacturer: str | None = None
    manufacturer_scores: dict[str, int] = field(default_factory=dict)
    products: list[str] = field(default_factory=list)
    fireproofing_type: str | None = None
    ratings: list[str] = field(default_factory=list)
    ul_designs: list[str] = field(default_factory=list)  # strict: "BXUV.X" or "Design No. X"
    ul_design_candidates: list[str] = field(default_factory=list)  # bare tokens like N708 (lower confidence)
    other_listings: list[str] = field(default_factory=list)  # Intertek/WH style ids
    member_terms: list[str] = field(default_factory=list)
    years_seen: list[int] = field(default_factory=list)
    dates_seen: list[str] = field(default_factory=list)
    project_hint: str | None = None
    current_or_historical: str = "historical"  # everything on the drive defaults to historical
    text_extracted: bool = False
    error: str | None = None


# ---------------------------------------------------------------------------------------------
def clean(s: str) -> str:
    """Make any str safe for UTF-8 outputs (undecodable filename bytes become U+FFFD)."""
    return s.encode("utf-8", "replace").decode("utf-8")


def iso_mtime(ts: float) -> str:
    try:
        return (datetime(1970, 1, 1, tzinfo=timezone.utc) + timedelta(seconds=ts)).astimezone().isoformat(
            timespec="seconds"
        )
    except (OverflowError, OSError, ValueError):
        return ""


def find_candidate_mounts() -> list[str]:
    cands: list[str] = []
    if os.name == "nt":
        import ctypes
        import string

        k32 = ctypes.windll.kernel32  # type: ignore[attr-defined]
        for letter in string.ascii_uppercase:
            d = f"{letter}:\\"
            if k32.GetDriveTypeW(d) not in (2, 3):  # DRIVE_REMOVABLE, DRIVE_FIXED
                continue
            buf = ctypes.create_unicode_buffer(261)
            ok = k32.GetVolumeInformationW(d, buf, 261, None, None, None, None, 0)
            cands.append(f"{d}  [{buf.value if ok else '?'}]")
    else:
        for r in ("/Volumes", "/media", "/mnt", "/run/media"):
            p = Path(r)
            if not p.is_dir():
                continue
            for child in sorted(p.iterdir()):
                try:
                    if child.is_dir():
                        cands.append(str(child))
                        if r in ("/media", "/run/media"):
                            cands.extend(str(c) for c in sorted(child.iterdir()) if c.is_dir())
                except OSError:
                    pass
    lacie = [c for c in cands if re.search(r"lac[ie]+y?", c, re.IGNORECASE)]
    return lacie + [c for c in cands if c not in lacie]


def same_volume(a: Path, b: Path) -> bool:
    if os.name == "nt":
        return os.path.splitdrive(str(a))[0].upper() == os.path.splitdrive(str(b))[0].upper()
    try:
        return a.stat().st_dev == (b if b.exists() else b.parent).stat().st_dev
    except OSError:
        return False


def is_reparse_or_link(path: str) -> bool:
    try:
        st = os.lstat(path)
    except OSError:
        return True
    if stat.S_ISLNK(st.st_mode):
        return True
    attrs = getattr(st, "st_file_attributes", 0)
    return bool(attrs & WIN_REPARSE)


def sha256_of(path: str, limit_mb: int) -> str | None:
    h = hashlib.sha256()
    try:
        with open(path, "rb") as f:
            read = 0
            while True:
                chunk = f.read(1 << 20)
                if not chunk:
                    break
                h.update(chunk)
                read += len(chunk)
                if limit_mb and read >= limit_mb << 20:
                    return "partial:" + h.hexdigest()
        return h.hexdigest()
    except OSError:
        return None


def read_pdf(path: str, e: Entry, max_pages: int, snippet_len: int) -> str:
    import pymupdf

    doc = pymupdf.open(path)  # read-only by path; nothing loaded into RAM beyond pages touched
    try:
        if doc.is_encrypted and not doc.authenticate(""):
            e.error = "encrypted PDF (password required)"
            e.pages = len(doc) if doc.page_count else None
            return ""
        e.pages = len(doc)
        meta = doc.metadata or {}
        e.title_meta = clean(meta.get("title") or "")[:200] or None
        e.author_meta = clean(meta.get("author") or "")[:120] or None
        e.created_meta = (meta.get("creationDate") or None) or None
        text_pages = 0
        img_pages = 0
        parts: list[str] = []
        for i in range(min(len(doc), max_pages)):
            try:
                pg = doc[i]
                t = pg.get_text("text") or ""
            except Exception as ex:  # noqa: BLE001
                parts.append("")
                e.error = f"page {i + 1}: {type(ex).__name__}"
                continue
            if t.strip():
                text_pages += 1
            try:
                if pg.get_images():
                    img_pages += 1
            except Exception:  # noqa: BLE001
                pass
            parts.append(t)
        e.has_text_layer = text_pages > 0
        e.likely_scanned = text_pages == 0 and img_pages > 0
        full = "\n".join(parts)
        e.snippet = clean(re.sub(r"\s+", " ", full)[:snippet_len])
        e.text_extracted = bool(full.strip())
        return full
    finally:
        doc.close()


def read_xlsx(path: str, e: Entry, snippet_len: int) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    try:
        e.pages = len(wb.sheetnames)
        e.title_meta = clean("; ".join(wb.sheetnames))[:200]
        for name in wb.sheetnames[:12]:
            try:
                ws = wb[name]
                parts.append(name)
                for r, row in enumerate(ws.iter_rows(values_only=True)):
                    if r >= 60:
                        break
                    parts.append(" ".join(str(v) for v in row if v is not None))
            except Exception as ex:  # noqa: BLE001 - chartsheets etc.
                parts.append(f"[sheet {name}: {type(ex).__name__}]")
    finally:
        wb.close()
    full = "\n".join(parts)
    e.snippet = clean(re.sub(r"\s+", " ", full)[:snippet_len])
    e.text_extracted = bool(full.strip())
    return full


def read_docx(path: str, e: Entry, snippet_len: int) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        e.error = "python-docx not installed; docx text skipped"
        return ""
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs[:400]]
    for t in d.tables[:10]:
        for row in t.rows[:40]:
            parts.append(" ".join(c.text for c in row.cells))
    full = "\n".join(parts)
    e.snippet = clean(re.sub(r"\s+", " ", full)[:snippet_len])
    e.text_extracted = bool(full.strip())
    return full


def _kw_hits(hay: str, kw: str) -> int:
    # word-ish boundaries so "ppg" does not match "ppgx" and "bid" does not match "forbidden"
    return len(re.findall(r"(?<![a-z0-9])" + re.escape(kw) + r"(?![a-z0-9])", hay))


def classify(e: Entry, text: str) -> None:
    text = clean(text)
    hay = re.sub(r"\s+", " ", (e.filename + " " + e.folder.replace(os.sep, " ") + " " + text).lower())
    fname = e.filename.lower()
    scores: dict[str, int] = {}
    for dt, kws in DOC_TYPES.items():
        s = 0
        for kw in kws:
            n = _kw_hits(hay, kw)
            if n:
                s += min(n, 5) + (3 if _kw_hits(fname, kw) else 0)
        if s:
            scores[dt] = s
    for rx, dt, bonus in FILENAME_RULES:
        if rx.search(fname):
            scores[dt] = scores.get(dt, 0) + bonus
    e.doc_type_scores = dict(sorted(scores.items(), key=lambda kv: -kv[1])[:5])
    if scores:
        best, bs = max(scores.items(), key=lambda kv: kv[1])
        second = sorted(scores.values(), reverse=True)[1] if len(scores) > 1 else 0
        conf = round(max(0.0, min(0.95, 0.3 + 0.1 * bs - 0.05 * second)), 2)
        if conf >= 0.45 and bs >= 2:
            e.doc_type, e.doc_type_confidence = best, conf
        else:
            e.doc_type, e.doc_type_confidence = "unclassified", conf
    mscores: dict[str, int] = {}
    for mf, kws in MANUFACTURERS.items():
        hits = {kw: _kw_hits(hay, kw) for kw in kws}
        tot = sum(hits.values())
        if tot:
            mscores[mf] = tot
            e.products.extend(k for k, n in hits.items() if n and k not in e.products)
    if mscores:
        e.manufacturer_scores = dict(sorted(mscores.items(), key=lambda kv: -kv[1]))
        e.manufacturer = max(mscores.items(), key=lambda kv: kv[1])[0]
    fscores = {ft: sum(_kw_hits(hay, kw) for kw in kws) for ft, kws in FP_TYPE.items()}
    fscores = {k: v for k, v in fscores.items() if v}
    if fscores:
        e.fireproofing_type = max(fscores.items(), key=lambda kv: kv[1])[0]
    scan = e.filename + "\n" + e.folder + "\n" + text
    e.ratings = sorted({re.sub(r"\s+", "", m.group("r")).replace("-", "-") + " HR" for m in RATING_RE.finditer(scan)})[:8]
    strict = {m.group("d").upper().replace("-", "") for m in UL_STRICT_RE.finditer(scan)}
    loose = {m.group("d").upper().replace("-", "") for m in UL_LOOSE_RE.finditer(scan)} - strict
    e.ul_designs = sorted(strict)[:30]
    e.ul_design_candidates = sorted(loose)[:30]
    e.other_listings = sorted({m.group(0).upper().replace(" ", "") for m in INTERTEK_RE.finditer(scan)})[:20]
    members = {m.group(1).upper().replace(" ", "") for m in MEMBER_RE.finditer(scan)}
    members |= {m.group(1).upper() for m in MEMBER_WORDS_RE.finditer(scan)}
    e.member_terms = sorted(members)[:25]
    e.years_seen = sorted({int(y) for y in YEAR_RE.findall(scan)})[:12]
    e.dates_seen = sorted({(a or b) for a, b in DATE_RE.findall(scan)})[:8]
    m = PROJECT_HINT_RE.search(scan)
    if m:
        e.project_hint = clean((m.group("p") or m.group("q") or "").strip())[:80] or None


# ---------------------------------------------------------------------------------------------
class Catalog:
    """Streams rows to catalog.jsonl as they are produced; builds the rest at the end."""

    def __init__(self, out: Path, root: Path, media_mode: str = "count", excludes: list[str] | None = None):
        self.out = out
        self.root = root
        self.entries: list[Entry] = []
        self.interrupted = False
        self.dirs_failed = 0
        self.media_mode = media_mode  # count | list
        self.excludes = [x.lower() for x in (excludes or [])]
        self.media: dict[str, dict[str, int]] = {}  # top folder -> {files, bytes, image, video, audio}
        self.excluded_dirs: list[str] = []
        out.mkdir(parents=True, exist_ok=True)
        # probe every output up front so a locked file fails fast, not after hours
        for name in ("catalog.jsonl", "catalog.csv", "catalog.sqlite", "summary.md", "summary_paths.md", "run.log"):
            p = out / name
            if p.exists():
                p.rename(p.with_name(name + ".prev"))
            with p.open("a", encoding="utf-8"):
                pass
        self._jsonl = (out / "catalog.jsonl").open("w", encoding="utf-8", errors="replace")
        self._log = (out / "run.log").open("w", encoding="utf-8", errors="replace")
        self.log(f"start root={root}")

    def count_media(self, folder: str, fclass: str, size: int) -> None:
        top = folder.split(os.sep)[0] if folder else "(root)"
        m = self.media.setdefault(top, {"files": 0, "bytes": 0, "image": 0, "video": 0, "audio": 0})
        m["files"] += 1
        m["bytes"] += size
        m[fclass] += 1

    def excluded(self, name: str) -> bool:
        import fnmatch

        n = name.lower()
        return any(fnmatch.fnmatch(n, pat) for pat in self.excludes)

    def log(self, msg: str) -> None:
        line = f"{datetime.now().isoformat(timespec='seconds')} {msg}"
        self._log.write(line + "\n")
        self._log.flush()
        print(line, file=sys.stderr)

    def add(self, e: Entry) -> None:
        self.entries.append(e)
        self._jsonl.write(json.dumps(asdict(e), ensure_ascii=False) + "\n")
        if len(self.entries) % 25 == 0:
            self._jsonl.flush()

    def close(self) -> None:
        self._jsonl.close()
        self._log.close()


def _entry_for_error(root: Path, path: str, msg: str, file_class: str = "other") -> Entry:
    p = Path(path)
    try:
        rel = p.relative_to(root)
    except ValueError:
        rel = p
    folder = "" if rel.parent == Path(".") else str(rel.parent)
    return Entry(clean(str(p)), clean(str(rel)), clean(folder), clean(p.name), p.suffix.lower(), 0, "", None,
                 file_class, error=clean(msg))


def inventory(root: Path, cat: Catalog, max_files: int | None, hash_limit_mb: int, max_pages: int,
              snippet_len: int, dry_run: bool) -> None:
    n = 0
    t0 = time.time()
    walk_root = str(root)
    if os.name == "nt" and not walk_root.startswith("\\\\?\\"):
        walk_root = "\\\\?\\" + walk_root  # lift the 260-char path limit
    prefix = "\\\\?\\" if walk_root.startswith("\\\\?\\") else ""

    def strip(p: str) -> str:
        return p[len(prefix):] if prefix and p.startswith(prefix) else p

    def on_err(ex: OSError) -> None:
        cat.dirs_failed += 1
        where = strip(str(getattr(ex, "filename", "") or "?"))
        cat.add(_entry_for_error(root, where, f"DIR UNREADABLE {type(ex).__name__}: {ex}", "directory"))

    try:
        for dirpath, dirnames, filenames in os.walk(walk_root, onerror=on_err):
            keep = []
            for d in dirnames:
                if d in SKIP_NAMES or d.startswith("."):
                    continue
                if cat.excluded(d):
                    cat.excluded_dirs.append(strip(os.path.join(dirpath, d)))
                    continue
                if is_reparse_or_link(os.path.join(dirpath, d)):
                    cat.add(_entry_for_error(root, strip(os.path.join(dirpath, d)), "link/junction not followed", "directory"))
                    continue
                keep.append(d)
            dirnames[:] = keep
            for fn in filenames:
                if fn in SKIP_NAMES or fn.startswith("._") or fn.startswith("~$"):
                    continue
                if max_files and n >= max_files:
                    cat.log(f"stopped at --max-files {max_files}")
                    return
                n += 1
                full = os.path.join(dirpath, fn)
                shown = strip(full)
                try:
                    _one_file(root, cat, full, shown, hash_limit_mb, max_pages, snippet_len, dry_run)
                except Exception as ex:  # noqa: BLE001 - never let one file end the run
                    cat.add(_entry_for_error(root, shown, f"UNHANDLED {type(ex).__name__}: {ex}"[:300]))
                if n % 50 == 0:
                    cat.log(f"{n} files, {time.time() - t0:.0f}s, last: {clean(os.path.basename(shown))[:60]}")
    except KeyboardInterrupt:
        cat.interrupted = True
        cat.log("INTERRUPTED by user — writing partial catalog")


def _one_file(root: Path, cat: Catalog, full: str, shown: str, hash_limit_mb: int, max_pages: int,
              snippet_len: int, dry_run: bool) -> None:
    p = Path(shown)
    if is_reparse_or_link(full):
        cat.add(_entry_for_error(root, shown, "symlink not followed"))
        return
    try:
        st = os.stat(full)
    except OSError as ex:
        cat.add(_entry_for_error(root, shown, f"STAT {type(ex).__name__}: {ex}"))
        return
    try:
        rel = p.relative_to(root)
    except ValueError:
        rel = p
    folder = "" if rel.parent == Path(".") else str(rel.parent)
    ext = p.suffix.lower()
    fclass = ("pdf" if ext in PDF_EXT else "excel" if ext in XLS_EXT else "word" if ext in DOC_EXT
              else "image" if ext in IMG_EXT else "video" if ext in VIDEO_EXT else "audio" if ext in AUDIO_EXT
              else "archive" if ext in ARCHIVE_EXT else "cad" if ext in CAD_EXT
              else "text" if ext in TEXT_EXT else "other")
    mt = iso_mtime(st.st_mtime)
    is_media = fclass in MEDIA_CLASSES
    if is_media and cat.media_mode == "count":
        cat.count_media(folder, fclass, st.st_size)
        return
    e = Entry(
        path=clean(shown), rel_path=clean(str(rel)), folder=clean(folder), filename=clean(p.name), ext=ext,
        size_bytes=st.st_size, mtime_iso=mt,
        sha256=None if (dry_run or is_media) else sha256_of(full, hash_limit_mb),
        file_class=fclass, doc_type=("media" if is_media else "unclassified"),
    )
    if not mt:
        e.error = f"bad mtime {st.st_mtime}"
    if is_media:
        cat.add(e)  # listed by name only; never opened
        return
    text = ""
    if not dry_run and st.st_size > 0:
        try:
            if fclass == "pdf":
                text = read_pdf(full, e, max_pages, snippet_len)
            elif fclass == "excel" and ext != ".xls":
                text = read_xlsx(full, e, snippet_len)
            elif fclass == "word" and ext == ".docx":
                text = read_docx(full, e, snippet_len)
            elif ext in (".xls", ".doc"):
                e.error = "legacy binary Office format; text not extracted (convert to .xlsx/.docx to index)"
            elif fclass == "text":
                with open(full, "rb") as f:
                    text = f.read(200_000).decode("utf-8", "replace")
                e.snippet = clean(re.sub(r"\s+", " ", text)[:snippet_len])
                e.text_extracted = bool(text.strip())
        except Exception as ex:  # noqa: BLE001 - record and continue
            e.error = f"{type(ex).__name__}: {ex}"[:300]
    elif st.st_size == 0:
        e.error = "zero-byte file"
    classify(e, text)
    cat.add(e)


# ---------------------------------------------------------------------------------------------
def write_outputs(cat: Catalog) -> None:
    out, root, entries = cat.out, cat.root, cat.entries
    cols = ["rel_path", "filename", "ext", "file_class", "size_bytes", "mtime_iso", "pages", "has_text_layer",
            "likely_scanned", "doc_type", "doc_type_confidence", "manufacturer", "fireproofing_type", "ratings",
            "ul_designs", "ul_design_candidates", "other_listings", "years_seen", "project_hint", "title_meta",
            "sha256", "error"]
    tmp = out / "catalog.csv.tmp"
    with tmp.open("w", newline="", encoding="utf-8-sig", errors="replace") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for e in entries:
            d = asdict(e)
            w.writerow([("; ".join(map(str, d[c])) if isinstance(d[c], list) else d[c]) for c in cols])
    os.replace(tmp, out / "catalog.csv")

    db_tmp = out / "catalog.sqlite.tmp"
    if db_tmp.exists():
        db_tmp.unlink()
    con = sqlite3.connect(db_tmp)
    fields = [f for f in Entry.__dataclass_fields__]
    con.execute("CREATE TABLE files (" + ", ".join(f"{f} TEXT" for f in fields) + ")")
    con.executemany(
        f"INSERT INTO files VALUES ({','.join('?' * len(fields))})",
        [tuple(json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v for v in asdict(e).values())
         for e in entries],
    )
    con.commit()
    con.close()
    os.replace(db_tmp, out / "catalog.sqlite")

    def count(key):
        c: dict[str, int] = {}
        for e in entries:
            v = key(e) or "(none)"
            c[v] = c.get(v, 0) + 1
        return sorted(c.items(), key=lambda kv: -kv[1])

    files = [e for e in entries if e.file_class != "directory"]
    status = "PARTIAL RUN (interrupted)" if cat.interrupted else "complete"
    volume = root.name or str(root)
    lines = [f"# LaCie library inventory — {datetime.now():%Y-%m-%d %H:%M} — {status}", "",
             f"Volume: `{volume}`  Files: {len(files)}  Total size: {sum(e.size_bytes for e in files) / 1e9:.2f} GB  "
             f"Unreadable directories: {cat.dirs_failed}", "",
             "## By file class", *[f"- {k}: {v}" for k, v in count(lambda e: e.file_class)], "",
             "## By document type (heuristic)", *[f"- {k}: {v}" for k, v in count(lambda e: e.doc_type)], "",
             "## By manufacturer (keyword hit)", *[f"- {k}: {v}" for k, v in count(lambda e: e.manufacturer)], "",
             "## By fireproofing type", *[f"- {k}: {v}" for k, v in count(lambda e: e.fireproofing_type)], "",
             "## By earliest year mentioned (decade)",
             *[f"- {k}: {v}" for k, v in sorted(count(lambda e: f"{e.years_seen[0] // 10 * 10}s" if e.years_seen else None))], "",
             "## Text extracted", f"- yes: {sum(1 for e in files if e.text_extracted)}",
             f"- no: {sum(1 for e in files if not e.text_extracted)}", "",
             "## Scanned PDFs (no text layer) — OCR candidates", f"- {sum(1 for e in files if e.likely_scanned)}", "",
             "## Media (photos/video/audio) — counted only, never opened",
             f"- files: {sum(m['files'] for m in cat.media.values())}  size: {sum(m['bytes'] for m in cat.media.values()) / 1e9:.1f} GB  "
             f"in {len(cat.media)} top-level folders" if cat.media else "- none counted (media mode: list)", "",
             "## Excluded folders", f"- {len(cat.excluded_dirs)}", "",
             "## Rows with errors", f"- {sum(1 for e in entries if e.error)} (see summary_paths.md)", "",
             "## Unclassified", f"- {sum(1 for e in files if e.doc_type == 'unclassified')} (see summary_paths.md)", ""]
    (out / "summary.md").write_text("\n".join(lines), encoding="utf-8", errors="replace")
    plines = ["# Inventory paths (KEEP LOCAL — contains folder and file names)", "", f"Root: `{root}`", "",
              "## Top folders", *[f"- {k}: {v}" for k, v in count(lambda e: e.folder.split(os.sep)[0] if e.folder else "(root)")[:60]], "",
              "## Media per top-level folder (counted only)",
              *[f"- {k}: {v['files']} files, {v['bytes'] / 1e9:.1f} GB (img {v['image']}, video {v['video']}, audio {v['audio']})"
                for k, v in sorted(cat.media.items(), key=lambda kv: -kv[1]['files'])], "",
              "## Excluded folders", *[f"- {d}" for d in cat.excluded_dirs[:200]], "",
              "## Errors (first 300)", *[f"- {e.rel_path}: {e.error}" for e in entries if e.error][:300], "",
              "## Unclassified (first 300)", *[f"- {e.rel_path}" for e in files if e.doc_type == "unclassified"][:300]]
    (out / "summary_paths.md").write_text("\n".join(plines), encoding="utf-8", errors="replace")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("root", nargs="?", help="library root (the LaCie drive or a folder on it)")
    ap.add_argument("--out", default=str(Path(__file__).resolve().parents[2] / "data" / "lacie_catalog"),
                    help="output directory (default: <repo>/data/lacie_catalog; must be on another volume)")
    ap.add_argument("--find", action="store_true", help="list candidate drives / mount points and exit")
    ap.add_argument("--max-files", type=int, default=None)
    ap.add_argument("--hash-limit-mb", type=int, default=16, help="hash only the first N MB of each file (0 = all)")
    ap.add_argument("--max-pages", type=int, default=4, help="PDF pages to read for classification")
    ap.add_argument("--snippet", type=int, default=600, help="max chars of text snippet stored per file")
    ap.add_argument("--dry-run", action="store_true", help="walk and stat only; do not open file contents")
    ap.add_argument("--allow-same-volume", action="store_true", help="(not recommended) permit --out on the root's volume")
    ap.add_argument("--media", choices=["count", "list"], default="count",
                    help="photos/video/audio: 'count' per folder (default) or 'list' every file by name (never opened)")
    ap.add_argument("--exclude", action="append", default=[], metavar="PATTERN",
                    help="folder name pattern to skip entirely (repeatable, case-insensitive glob), e.g. --exclude 'Photos*'")
    a = ap.parse_args()
    if a.find:
        for c in find_candidate_mounts():
            print(c)
        return 0
    if not a.root:
        ap.error("root is required (or use --find)")
    root_s = a.root.strip().strip('"').strip("'")
    if os.name == "nt" and re.fullmatch(r"[A-Za-z]:", root_s):
        root_s += "\\"
    root = Path(root_s).resolve()
    out = Path(a.out).resolve()
    if not root.is_dir():
        print(f"root is not a directory: {root}", file=sys.stderr)
        return 2
    rn, on = os.path.normcase(str(root)), os.path.normcase(str(out))
    if on == rn or on.startswith(rn.rstrip("\\/") + os.sep):
        print(f"refusing to write output inside the library root: {out}", file=sys.stderr)
        return 2
    if same_volume(root, out) and not a.allow_same_volume:
        print(f"refusing: output {out} is on the same volume as the library root {root}. "
              f"Use --out on your internal drive.", file=sys.stderr)
        return 2
    print(f"inventorying {root} (read-only) → {out}", file=sys.stderr)
    cat = Catalog(out, root, media_mode=a.media, excludes=a.exclude)
    try:
        inventory(root, cat, a.max_files, a.hash_limit_mb, a.max_pages, a.snippet, a.dry_run)
    finally:
        cat.close()
    if not cat.entries and not cat.media:
        print("No files found. Check the path with --find. On macOS, allow Terminal under System Settings > "
              "Privacy & Security > Files and Folders > Removable Volumes.", file=sys.stderr)
        return 2
    write_outputs(cat)
    status = "PARTIAL" if cat.interrupted else "done"
    print(f"{status}: {len(cat.entries)} rows ({cat.dirs_failed} unreadable directories). See {cat.out / 'summary.md'}",
          file=sys.stderr)
    return 3 if cat.interrupted else 0


if __name__ == "__main__":
    sys.exit(main())
